"""
Geometry QA on the real file, since LibreOffice cannot render anything in this
sandbox. Catches the two defects that actually reach a viewer: content placed
outside the canvas, and text that needs more room than its box has.
"""
from pptx import Presentation
from pptx.util import Emu
import math, sys

EMU_IN = 914400.0
W, H = 13.3, 7.5
MARGIN = 0.4                      # slide edge tolerance

# Rough advance width as a fraction of point size. Cambria is a touch wider.
ADV = {'Cambria': 0.50, 'Calibri': 0.47}

def inches(v): return (v or 0) / EMU_IN

def est_lines(text, font_pt, box_w_in, face):
    if not text.strip(): return 0
    per_char = ADV.get(face, 0.5) * font_pt / 72.0     # inches
    if per_char <= 0: return 1
    cpl = max(1, int(box_w_in / per_char))
    lines = 0
    for para in text.split('\n'):
        lines += max(1, math.ceil(len(para) / cpl))
    return lines

prs = Presentation(sys.argv[1])
problems = []

for n, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        x, y = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        boxes.append((sh, x, y, w, h))

        # 1. off-canvas
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            problems.append(f"slide {n}: OFF-CANVAS  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}  "
                            f"({(sh.text_frame.text[:32] if sh.has_text_frame else sh.shape_type)})")
        # 2. too close to the edge
        elif x < MARGIN - 0.01 or y < MARGIN - 0.01 or x + w > W - MARGIN + 0.01 or y + h > H - MARGIN + 0.01:
            label = sh.text_frame.text[:28].replace('\n', ' ') if sh.has_text_frame else str(sh.shape_type)
            problems.append(f"slide {n}: TIGHT MARGIN x={x:.2f} y={y:.2f} r={x+w:.2f} b={y+h:.2f}  ({label})")

        # 3. text that needs more height than the box has
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            runs = [r for p in tf.paragraphs for r in p.runs]
            if not runs: continue
            pt = max((r.font.size.pt if r.font.size else 18) for r in runs)
            face = next((r.font.name for r in runs if r.font.name), 'Calibri')
            lines = est_lines(tf.text, pt, w - 0.1, face)
            need = lines * pt * 1.22 / 72.0
            if need > h + 0.06:
                problems.append(f"slide {n}: OVERFLOW?   needs ~{need:.2f}\" in {h:.2f}\"  "
                                f"{pt:.0f}pt x{lines}ln  ({tf.text[:34].replace(chr(10),' ')})")

print(f"{len(prs.slides)} slides checked")
print("\n".join(problems) if problems else "no geometry problems found")
